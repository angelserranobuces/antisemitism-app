"""Report generation: XLSX, CSV, JSON, DOCX outputs."""

import csv
import io
import json
from datetime import date
from typing import List, Dict, Any, Optional

import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter

from app.analyzer import CATEGORIES

# ── Helpers ───────────────────────────────────────────────────────────────────

CATEGORY_LABELS = [meta["label"] for meta in CATEGORIES.values()]
CATEGORY_KEYS = list(CATEGORIES.keys())

RISK_COLOR = {
    "No Risk":      "FFFFFF",
    "Very Low":     "E8F5E9",
    "Low":          "C8E6C9",
    "Low-Moderate": "FFF9C4",
    "Moderate":     "FFE0B2",
    "High":         "FFCDD2",
    "Very High":    "B71C1C",
    "Error":        "F5F5F5",
}

def score_to_risk_level(score: float) -> str:
    if score < 0.3:   return "No Risk"
    if score < 0.8:   return "Very Low"
    if score < 1.3:   return "Low"
    if score < 1.9:   return "Low-Moderate"
    if score < 2.8:   return "Moderate"
    if score < 3.8:   return "High"
    return "Very High"

def _header_style(ws, cell_ref, fill_hex="1565C0"):
    cell = ws[cell_ref]
    cell.font = Font(bold=True, color="FFFFFF", size=11)
    cell.fill = PatternFill("solid", fgColor=fill_hex)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def _thin_border():
    s = Side(style="thin")
    return Border(left=s, right=s, top=s, bottom=s)


# ── 1. Per-Newspaper JSON ─────────────────────────────────────────────────────

def generate_newspaper_json(
    newspaper: str,
    analysis_date: str,
    analysis: Dict[str, Any],
    filename: str,
) -> bytes:
    payload = {
        "newspaper": newspaper,
        "date": analysis_date,
        "source_filename": filename,
        "methodology_version": analysis.get("methodology_version", "1.0"),
        "processing_timestamp": analysis.get("processing_timestamp", ""),
        "category_scores": analysis["category_scores"],
        "overall_rating": analysis["overall_rating"],
        "risk_level": analysis.get("risk_level", score_to_risk_level(analysis["overall_rating"])),
        "confidence_overall": analysis.get("confidence_overall", ""),
        "main_drivers": analysis.get("main_drivers", ""),
        "mitigating_factors": analysis.get("mitigating_factors", ""),
        "evidence": analysis.get("evidence_items", []),
        "framework_references": analysis.get("framework_references", []),
    }
    return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")


# ── 2. Daily Excel Dataset ────────────────────────────────────────────────────

def generate_daily_excel(results: List[Dict], analysis_date: str) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Daily Dataset"

    headers = ["Date", "Newspaper", "Overall_Rating", "Risk_Level"] + CATEGORY_LABELS + ["Source_Filename"]
    ws.append(headers)

    # Style headers
    for col_idx, _ in enumerate(headers, 1):
        _header_style(ws, f"{get_column_letter(col_idx)}1")

    for r in results:
        scores = r.get("category_scores", {})
        row = [
            r.get("date", analysis_date),
            r.get("newspaper", ""),
            r.get("overall_rating", 0),
            score_to_risk_level(r.get("overall_rating", 0)),
        ]
        for key in CATEGORY_KEYS:
            s = scores.get(key, {})
            score_val = s.get("score", 0) if isinstance(s, dict) else s
            row.append(score_val if score_val != "Not Determined" else "N/D")
        row.append(r.get("source_filename", ""))
        ws.append(row)

    # Auto-width
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 40)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── 3. Ranking Table ──────────────────────────────────────────────────────────

def generate_ranking_excel(results: List[Dict]) -> bytes:
    ranked = sorted(results, key=lambda x: x.get("overall_rating", 0), reverse=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ranking"

    headers = ["Rank", "Newspaper", "Overall_Rating", "Risk_Level", "Main_Drivers"]
    ws.append(headers)
    for col_idx, _ in enumerate(headers, 1):
        _header_style(ws, f"{get_column_letter(col_idx)}1")

    for i, r in enumerate(ranked, 1):
        overall = r.get("overall_rating", 0)
        risk = score_to_risk_level(overall)
        ws.append([
            i,
            r.get("newspaper", ""),
            overall,
            risk,
            r.get("main_drivers", ""),
        ])
        # Color-code risk level cell
        color = RISK_COLOR.get(risk, "FFFFFF")
        ws.cell(row=i + 1, column=4).fill = PatternFill("solid", fgColor=color)

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 60)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_ranking_csv(results: List[Dict]) -> bytes:
    ranked = sorted(results, key=lambda x: x.get("overall_rating", 0), reverse=True)
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["Rank", "Newspaper", "Overall_Rating", "Risk_Level", "Main_Drivers"])
    for i, r in enumerate(ranked, 1):
        overall = r.get("overall_rating", 0)
        writer.writerow([
            i,
            r.get("newspaper", ""),
            overall,
            score_to_risk_level(overall),
            r.get("main_drivers", ""),
        ])
    return buf.getvalue().encode("utf-8-sig")


# ── 4. Expanded Evidence Table ────────────────────────────────────────────────

def _flatten_evidence(results: List[Dict]) -> List[Dict]:
    rows = []
    for r in results:
        newspaper = r.get("newspaper", "")
        scores = r.get("category_scores", {})
        for item in r.get("evidence_items", []):
            cat_key = item.get("category", "")
            cat_score = scores.get(cat_key, {})
            score_val = cat_score.get("score", "N/A") if isinstance(cat_score, dict) else "N/A"
            cat_label = CATEGORIES.get(cat_key, {}).get("label", cat_key)
            rows.append({
                "Newspaper": newspaper,
                "Page_Number": item.get("page_reference", "N/A"),
                "Article_Title": item.get("article_title", "N/A"),
                "Category": cat_label,
                "Trigger_Phrase": item.get("trigger_phrase", ""),
                "Evidence_Notes": item.get("reason", ""),
                "Category_Score": score_val,
                "Impact": item.get("impact", ""),
            })
    return rows


def generate_evidence_excel(results: List[Dict]) -> bytes:
    rows = _flatten_evidence(results)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Evidence"

    headers = ["Newspaper", "Page_Number(s)", "Article_Title", "Category",
               "Trigger_Phrase", "Evidence_Notes", "Related_Category_Score", "Impact"]
    ws.append(headers)
    for col_idx, _ in enumerate(headers, 1):
        _header_style(ws, f"{get_column_letter(col_idx)}1", fill_hex="4A148C")

    for row in rows:
        ws.append([
            row["Newspaper"], row["Page_Number"], row["Article_Title"],
            row["Category"], row["Trigger_Phrase"], row["Evidence_Notes"],
            row["Category_Score"], row["Impact"],
        ])

    # Wrap text in evidence notes column
    for cell in ws["F"][1:]:
        cell.alignment = Alignment(wrap_text=True)

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col_letter].width = min(max_len + 4, 60)
    ws.row_dimensions[1].height = 30

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_evidence_csv(results: List[Dict]) -> bytes:
    rows = _flatten_evidence(results)
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=[
        "Newspaper", "Page_Number", "Article_Title", "Category",
        "Trigger_Phrase", "Evidence_Notes", "Category_Score", "Impact",
    ])
    writer.writeheader()
    writer.writerows(rows)
    return buf.getvalue().encode("utf-8-sig")


def generate_evidence_json(results: List[Dict]) -> bytes:
    rows = _flatten_evidence(results)
    return json.dumps(rows, ensure_ascii=False, indent=2).encode("utf-8")


# ── DOCX helpers ──────────────────────────────────────────────────────────────

def _cell_set_text(cell, text: str, bold: bool = False, font_size_pt: float = 9,
                   white_text: bool = False, fill_rgb: str = None):
    """
    Safely write text into a table cell and apply formatting.
    Uses paragraph.clear() + add_run() to avoid the runs[0] IndexError.
    """
    from docx.shared import Pt, RGBColor
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    para = cell.paragraphs[0]
    para.clear()                        # remove any existing runs
    run = para.add_run(str(text))       # always creates exactly one run
    run.bold = bold
    run.font.size = Pt(font_size_pt)
    if white_text:
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    if fill_rgb:
        tc   = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd  = OxmlElement("w:shd")
        shd.set(qn("w:val"),   "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"),  fill_rgb)
        tcPr.append(shd)


def _docx_table_hdr(table, row_idx: int, texts: List[str],
                    bold: bool = True, fill_rgb: str = None):
    """Style a header row: white bold text on a coloured background."""
    row = table.rows[row_idx]
    for i, text in enumerate(texts):
        _cell_set_text(row.cells[i], text, bold=bold, font_size_pt=9,
                       white_text=bool(fill_rgb), fill_rgb=fill_rgb)


def _docx_add_row(table, values: List[str], font_size_pt: float = 9):
    """Append a data row to a python-docx Table."""
    row = table.add_row()
    for i, val in enumerate(values):
        _cell_set_text(row.cells[i], str(val) if val is not None else "",
                       font_size_pt=font_size_pt)


def _docx_bytes(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── Report 1: Evidence Phrases & FRS Formula Walk-Through ────────────────────

CATEGORY_CODES = {
    "collective_blame":                 "CB",
    "conspiracy_tropes":                "CT",
    "holocaust_inversion":              "HI",
    "demonisation":                     "D",
    "denial_jewish_self_determination": "DSD",
    "double_standards":                 "DS",
    "dehumanisation":                   "DH",
    "asymmetric_empathy":               "AE",
    "contextual_omission":              "CO",
}


def generate_evidence_phrases_docx(results: List[Dict], analysis_date: str) -> bytes:
    """
    Generates 'Evidence Phrases and Formula Walk-Through' .docx report
    matching the reference document structure.
    """
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # ── Page margins ──────────────────────────────────────────────────────────
    for section in doc.sections:
        section.top_margin    = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin   = Inches(1)
        section.right_margin  = Inches(1)

    # ── Title ─────────────────────────────────────────────────────────────────
    title = doc.add_heading("Evidence Phrases and Formula Walk-Through", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph(f"Specific phrases driving the Framing Risk Score — {analysis_date}")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.size = Pt(11)
    sub.runs[0].italic = True

    doc.add_paragraph()

    # ── Section 1: Weighted formula table ─────────────────────────────────────
    doc.add_heading("Weighted formula", level=2)
    tbl = doc.add_table(rows=1, cols=3)
    tbl.style = "Table Grid"
    _docx_table_hdr(tbl, 0, ["Code", "Category", "Weight"], fill_rgb="2E75B6")

    for key, meta in CATEGORIES.items():
        code = CATEGORY_CODES[key]
        _docx_add_row(tbl, [code, meta["label"], f"{meta['weight']:.2f}"])

    doc.add_paragraph()

    # ── Section 2: Formula totals by newspaper ────────────────────────────────
    doc.add_heading("Formula totals by newspaper", level=2)

    sorted_results = sorted(results, key=lambda r: r.get("overall_rating", 0), reverse=True)

    tbl2 = doc.add_table(rows=1, cols=3)
    tbl2.style = "Table Grid"
    _docx_table_hdr(tbl2, 0, ["Newspaper", "Add-up through formula", "FRS"], fill_rgb="2E75B6")

    for r in sorted_results:
        scores = r.get("category_scores", {})
        parts = []
        for key, meta in CATEGORIES.items():
            sc = scores.get(key, {})
            raw = sc.get("score", 0) if isinstance(sc, dict) else sc
            try:
                raw = float(raw)
            except (ValueError, TypeError):
                raw = 0.0
            if raw > 0:
                code = CATEGORY_CODES[key]
                parts.append(f"{code} {raw}×{meta['weight']}={raw*meta['weight']:.3f}")
        formula_str = " + ".join(parts) if parts else "No indicators found"
        frs = r.get("overall_rating", 0)
        _docx_add_row(tbl2, [r.get("newspaper", ""), formula_str, f"{frs:.3f}"])

    doc.add_paragraph()

    # ── Section 3: Phrase-by-phrase scoring rationale ─────────────────────────
    doc.add_heading("Phrase-by-phrase scoring rationale", level=2)

    tbl3 = doc.add_table(rows=1, cols=5)
    tbl3.style = "Table Grid"
    _docx_table_hdr(tbl3, 0,
                    ["Newspaper", "Page", "Phrase", "Scoring effect", "Why it matters"],
                    fill_rgb="4A148C")

    for r in sorted_results:
        newspaper = r.get("newspaper", "")
        evidence = r.get("evidence_items", [])
        if not evidence:
            _docx_add_row(tbl3, [newspaper, "N/A", "No relevant phrases identified", "No scoring", "No indicators found"])
        for ev in evidence:
            phrase = ev.get("trigger_phrase", "")
            page   = ev.get("page_reference", "N/A")
            reason = ev.get("reason", "")
            cat    = ev.get("category", "")
            score_info = ""
            if cat in CATEGORIES and cat in (r.get("category_scores") or {}):
                sc = r["category_scores"][cat]
                raw = sc.get("score", 0) if isinstance(sc, dict) else sc
                code = CATEGORY_CODES.get(cat, cat)
                score_info = f"{code} {raw} → triggers this category"
            _docx_add_row(tbl3, [newspaper, str(page), phrase, score_info, reason])

    doc.add_paragraph()

    # ── Section 4: Category inputs by newspaper ───────────────────────────────
    doc.add_heading("Category inputs by newspaper", level=2)

    for r in sorted_results:
        newspaper = r.get("newspaper", "")
        scores    = r.get("category_scores", {})

        doc.add_heading(newspaper, level=3)

        tbl4 = doc.add_table(rows=1, cols=4)
        tbl4.style = "Table Grid"
        _docx_table_hdr(tbl4, 0, ["Code", "Category", "Raw score", "Weighted contribution"],
                        fill_rgb="1565C0")

        total_frs = 0.0
        for key, meta in CATEGORIES.items():
            code = CATEGORY_CODES[key]
            sc   = scores.get(key, {})
            raw  = sc.get("score", 0) if isinstance(sc, dict) else sc
            try:
                raw = float(raw)
            except (ValueError, TypeError):
                raw = 0.0
            weighted = raw * meta["weight"]
            total_frs += weighted
            _docx_add_row(tbl4, [code, meta["label"], f"{raw:.1f}", f"{weighted:.3f}"])

        frs_val = r.get("overall_rating", total_frs)
        p = doc.add_paragraph(f"Total FRS for {newspaper}: {frs_val:.3f} / 5")
        p.runs[0].bold = True
        doc.add_paragraph()

    return _docx_bytes(doc)


# ── Report 2: Comprehensive Antisemitic Framing Risk Review ──────────────────

def generate_comprehensive_risk_docx(results: List[Dict], analysis_date: str) -> bytes:
    """
    Generates 'Comprehensive Antisemitic Framing Risk Review' .docx report
    matching the reference document structure.
    """
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    for section in doc.sections:
        section.top_margin    = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin   = Inches(1)
        section.right_margin  = Inches(1)

    # ── Title ─────────────────────────────────────────────────────────────────
    title = doc.add_heading("Comprehensive Antisemitic Framing Risk Review", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    newspaper_names = ", ".join(r.get("newspaper", "") for r in results)
    sub = doc.add_paragraph(f"{newspaper_names} — {analysis_date}")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].italic = True
    sub.runs[0].font.size = Pt(11)

    doc.add_paragraph()

    sorted_results = sorted(results, key=lambda r: r.get("overall_rating", 0), reverse=True)

    # ── Executive Summary ─────────────────────────────────────────────────────
    doc.add_heading("Executive summary", level=2)

    # Build summary narrative
    all_low = all(r.get("overall_rating", 0) < 2.0 for r in results)
    risk_summary = "a low-risk assessment across all newspapers" if all_low else "varied risk levels across newspapers"

    summary_text = (
        f"This report applies the Framing Risk Score (FRS) framework to the uploaded editions "
        f"of {newspaper_names} for {analysis_date}. The framework measures indicators of "
        f"antisemitic framing risk; it does not determine intent. The result is {risk_summary}."
    )
    doc.add_paragraph(summary_text)

    drivers_found = [r.get("main_drivers", "") for r in results if r.get("main_drivers")]
    if drivers_found:
        doc.add_paragraph("Key findings: " + " | ".join(d for d in drivers_found if d))

    doc.add_paragraph()

    # Ranking summary table
    tbl_rank = doc.add_table(rows=1, cols=4)
    tbl_rank.style = "Table Grid"
    _docx_table_hdr(tbl_rank, 0, ["Rank", "Newspaper", "FRS / 5", "Risk level"],
                    fill_rgb="2E75B6")

    for rank, r in enumerate(sorted_results, 1):
        frs  = r.get("overall_rating", 0)
        risk = r.get("risk_level", score_to_risk_level(frs))
        _docx_add_row(tbl_rank, [str(rank), r.get("newspaper", ""), f"{frs:.2f}", risk])

    doc.add_paragraph()

    # ── Source coverage ───────────────────────────────────────────────────────
    doc.add_heading("Source coverage and confidence", level=2)

    tbl_src = doc.add_table(rows=1, cols=4)
    tbl_src.style = "Table Grid"
    _docx_table_hdr(tbl_src, 0, ["Newspaper", "Pages", "Coverage note", "Confidence"],
                    fill_rgb="1565C0")

    for r in sorted_results:
        newspaper = r.get("newspaper", "")
        pages     = str(r.get("total_pages", "N/A"))
        mode      = r.get("extraction_mode", "text")
        conf      = r.get("confidence_overall", "Moderate")
        note      = ("Searchable text layer; full keyword pass performed."
                     if mode == "text"
                     else "Image-only PDF; OCR extraction applied.")
        _docx_add_row(tbl_src, [newspaper, pages, note, conf])

    doc.add_paragraph()

    # ── Detailed category scoring ─────────────────────────────────────────────
    doc.add_heading("Detailed category scoring", level=2)

    cat_headers = ["Newspaper"] + [meta["label"] for meta in CATEGORIES.values()] + ["FRS"]
    tbl_cat = doc.add_table(rows=1, cols=len(cat_headers))
    tbl_cat.style = "Table Grid"
    _docx_table_hdr(tbl_cat, 0, cat_headers, fill_rgb="2E75B6")

    for r in sorted_results:
        scores = r.get("category_scores", {})
        row_vals = [r.get("newspaper", "")]
        for key in CATEGORIES:
            sc  = scores.get(key, {})
            raw = sc.get("score", 0) if isinstance(sc, dict) else sc
            try:
                raw = float(raw)
            except (ValueError, TypeError):
                raw = 0.0
            row_vals.append(f"{raw:.1f}")
        row_vals.append(f"{r.get('overall_rating', 0):.3f}")
        _docx_add_row(tbl_cat, row_vals, font_size_pt=8)

    doc.add_paragraph()

    # ── Interpretation by newspaper ───────────────────────────────────────────
    doc.add_heading("Interpretation by newspaper", level=2)

    for r in sorted_results:
        newspaper = r.get("newspaper", "")
        doc.add_heading(newspaper, level=3)
        drivers    = r.get("main_drivers", "No significant drivers identified.")
        mitigating = r.get("mitigating_factors", "")
        doc.add_paragraph(drivers or "No significant drivers identified.")
        if mitigating:
            p = doc.add_paragraph("Mitigating factors: " + mitigating)
            p.runs[0].italic = True
        doc.add_paragraph()

    # ── Evidence table ────────────────────────────────────────────────────────
    doc.add_heading("Evidence table", level=2)

    tbl_ev = doc.add_table(rows=1, cols=5)
    tbl_ev.style = "Table Grid"
    _docx_table_hdr(tbl_ev, 0,
                    ["Newspaper", "Page", "Article / location", "Phrase / evidence", "Assessment"],
                    fill_rgb="4A148C")

    for r in sorted_results:
        newspaper = r.get("newspaper", "")
        evidence  = r.get("evidence_items", [])
        if not evidence:
            _docx_add_row(tbl_ev, [newspaper, "N/A", "N/A", "No indicators found", "N/A"])
        for ev in evidence:
            _docx_add_row(tbl_ev, [
                newspaper,
                str(ev.get("page_reference", "N/A")),
                ev.get("article_title", "N/A"),
                ev.get("trigger_phrase", ""),
                ev.get("reason", ""),
            ])

    doc.add_paragraph()

    # ── Methodological appendix ───────────────────────────────────────────────
    doc.add_heading("Methodological appendix", level=2)

    formula_str = " + ".join(
        f"({CATEGORY_CODES[k]}×{meta['weight']:.2f})"
        for k, meta in CATEGORIES.items()
    )
    doc.add_paragraph(
        f"Formula: FRS = [{formula_str}]. "
        f"Scale: 0 = no evidence; 1 = weak isolated indicator; 2 = recurrent low-intensity; "
        f"3 = clear repeated pattern; 4 = strong persistent pattern; 5 = extreme/dominant framing. "
        f"Caveat: the framework measures framing-risk indicators, not proven antisemitic intent. "
        f"Human review remains essential."
    )

    return _docx_bytes(doc)
